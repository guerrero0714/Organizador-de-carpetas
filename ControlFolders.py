"""
OrderFolder v4
PDF · Imagen · Word · Excel · Google Sheets
Vencimientos · Exportar · Tags · Modo claro/oscuro

Creado por Hack14

Requisitos:
    pip install matplotlib pillow pymupdf python-docx openpyxl reportlab

Uso: python orderfolder.py
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import json, os, subprocess, sys, shutil, webbrowser
from datetime import datetime, date, timedelta
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from collections import Counter
from PIL import Image, ImageTk

try:
    import fitz;       PDF_OK  = True
except ImportError:    PDF_OK  = False
try:
    import docx;       DOCX_OK = True
except ImportError:    DOCX_OK = False
try:
    import openpyxl;   XLSX_OK = True
except ImportError:    XLSX_OK = False
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors as rl_colors
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                    Paragraph, Spacer)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    RL_OK = True
except ImportError:    RL_OK   = False

# ──────────────────────────────────────────────
#  CONFIGURACIÓN
# ──────────────────────────────────────────────
APP_DIR   = os.path.dirname(os.path.abspath(__file__))
CONFIG_F  = os.path.join(APP_DIR, "config.json")
DATA_FILE = os.path.join(APP_DIR, "certificaciones.json")
EXTS_IMG  = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
EXTS_ALL  = {".pdf",".docx",".doc",".xlsx",".xls"} | EXTS_IMG
ALERT_DAYS = 30

TIPO_ICON = {".pdf":"📄",".png":"🖼",".jpg":"🖼",".jpeg":"🖼",
             ".webp":"🖼",".bmp":"🖼",".docx":"📝",".doc":"📝",
             ".xlsx":"📊",".xls":"📊","url":"🔗"}

# ── Temas ────────────────────────────────────
THEMES = {
    "dark": {
        "bg":"#080b14","bg2":"#0d1117","sidebar":"#090d16",
        "card":"#111827","card2":"#1a2235","card3":"#1e2a3a",
        "accent":"#4f46e5","cyan":"#06b6d4","green":"#10b981",
        "red":"#ef4444","yellow":"#f59e0b","purple":"#8b5cf6",
        "orange":"#f97316","text":"#f1f5f9","text2":"#cbd5e1",
        "sub":"#64748b","border":"#1e293b","border2":"#263044",
        "nav_sel":"#1e1b4b","tag_bg":"#1e2a3a","tag_fg":"#06b6d4",
    },
    "light": {
        "bg":"#f8fafc","bg2":"#f1f5f9","sidebar":"#e2e8f0",
        "card":"#ffffff","card2":"#f8fafc","card3":"#e2e8f0",
        "accent":"#4f46e5","cyan":"#0891b2","green":"#059669",
        "red":"#dc2626","yellow":"#d97706","purple":"#7c3aed",
        "orange":"#ea580c","text":"#0f172a","text2":"#1e293b",
        "sub":"#64748b","border":"#cbd5e1","border2":"#94a3b8",
        "nav_sel":"#e0e7ff","tag_bg":"#e0e7ff","tag_fg":"#4f46e5",
    }
}
C = dict(THEMES["dark"])   # tema activo (mutable)

FONT = "Segoe UI"
def fb(s): return (FONT, s, "bold")
def fn(s): return (FONT, s)

# ──────────────────────────────────────────────
#  PERSISTENCIA
# ──────────────────────────────────────────────
def load_cfg():
    try:
        if os.path.exists(CONFIG_F):
            with open(CONFIG_F,"r",encoding="utf-8") as f:
                return json.load(f)
    except Exception: pass
    return {"carpeta_raiz":"","theme":"dark"}

def save_cfg(c):
    with open(CONFIG_F,"w",encoding="utf-8") as f:
        json.dump(c,f,ensure_ascii=False,indent=2)

def load_data():
    try:
        if os.path.exists(DATA_FILE):
            with open(DATA_FILE,"r",encoding="utf-8") as f:
                return json.load(f)
    except Exception: pass
    return []

def save_data(d):
    with open(DATA_FILE,"w",encoding="utf-8") as f:
        json.dump(d,f,ensure_ascii=False,indent=2)

# ──────────────────────────────────────────────
#  ESCANEO
# ──────────────────────────────────────────────
def scan(raiz):
    out = []
    if not os.path.isdir(raiz): return out
    def add(path, cat):
        ext = os.path.splitext(path)[1].lower()
        out.append({"nombre":os.path.splitext(os.path.basename(path))[0],
                    "categoria":cat,"fecha":"","fecha_venc":"","descripcion":"",
                    "ruta":path,"archivo":"","tipo":ext,"url":"","tags":[]})
    try:
        for f in sorted(os.listdir(raiz)):
            p = os.path.join(raiz,f)
            if os.path.isfile(p) and os.path.splitext(f)[1].lower() in EXTS_ALL:
                add(p,"General")
        for sub in sorted(os.listdir(raiz)):
            sp = os.path.join(raiz,sub)
            if os.path.isdir(sp):
                for f in sorted(os.listdir(sp)):
                    p = os.path.join(sp,f)
                    if os.path.isfile(p) and os.path.splitext(f)[1].lower() in EXTS_ALL:
                        add(p,sub)
    except PermissionError: pass
    return out

def sync(existing, scanned):
    by_ruta = {c.get("ruta",""):c for c in existing}
    max_id  = max((c.get("id",0) for c in existing),default=0)
    result, seen = [], set()
    for item in scanned:
        r = item["ruta"]
        if r in by_ruta:
            orig = dict(by_ruta[r]); orig["ruta"]=r
            if orig.get("id") not in seen:
                result.append(orig); seen.add(orig.get("id"))
        else:
            max_id+=1; item["id"]=max_id
            result.append(item); seen.add(max_id)
    scanned_r = {i["ruta"] for i in scanned}
    for c in existing:
        if c.get("ruta","") not in scanned_r and c.get("id") not in seen:
            result.append(c); seen.add(c.get("id"))
    return result

# ──────────────────────────────────────────────
#  VENCIMIENTOS
# ──────────────────────────────────────────────
def days_to_expiry(fecha_venc_str):
    if not fecha_venc_str or not fecha_venc_str.strip():
        return None
    try:
        fv = datetime.strptime(fecha_venc_str.strip(), "%Y-%m-%d").date()
        return (fv - date.today()).days
    except ValueError:
        return None

def expiry_status(days):
    if days is None:   return "none",   ""
    if days < 0:       return "expired", f"Venció hace {abs(days)}d"
    if days <= ALERT_DAYS: return "soon", f"Vence en {days}d"
    return "ok", f"Vence en {days}d"

# ──────────────────────────────────────────────
#  LECTURA DOCX / XLSX
# ──────────────────────────────────────────────
def read_docx(path):
    if not DOCX_OK: return None,"pip install python-docx"
    try:
        doc = docx.Document(path)
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(lines[:120]) or "(vacío)", None
    except Exception as e: return None,str(e)

def read_xlsx(path):
    if not XLSX_OK: return None,"pip install openpyxl"
    try:
        wb = openpyxl.load_workbook(path,read_only=True,data_only=True)
        out = []
        for sname in wb.sheetnames[:3]:
            ws = wb[sname]; out.append(f"── {sname} ──")
            for row in ws.iter_rows(max_row=40,values_only=True):
                vals=[str(v) if v is not None else "" for v in row]
                if any(v.strip() for v in vals):
                    out.append("  │  ".join(vals[:8]))
            out.append("")
        wb.close()
        return "\n".join(out) or "(vacío)", None
    except Exception as e: return None,str(e)

def read_txt(path):
    try:
        with open(path,"r",encoding="utf-8",errors="replace") as f:
            lines = f.readlines()[:200]
        # Párrafos: separar bloques por líneas vacías
        out, block = [], []
        for line in lines:
            if line.strip():
                block.append(line.rstrip())
            else:
                if block:
                    out.append("\n".join(block))
                    out.append("")
                    block = []
        if block:
            out.append("\n".join(block))
        return "\n".join(out) or "(archivo vacío)", None
    except Exception as e: return None, str(e)

def read_csv(path):
    import csv
    try:
        with open(path,"r",encoding="utf-8",errors="replace") as f:
            sample = f.read(4096)
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        with open(path,"r",encoding="utf-8",errors="replace") as f:
            reader = csv.reader(f, dialect)
            rows   = [row for _, row in zip(range(60), reader)]
        if not rows: return "(CSV vacío)", None
        # Calcular ancho de columnas
        col_w = [max(len(str(r[c])) if c<len(r) else 0
                     for r in rows) for c in range(len(rows[0]))]
        col_w = [min(w, 24) for w in col_w]
        lines = []
        for i, row in enumerate(rows):
            padded = [str(v)[:col_w[c]].ljust(col_w[c])
                      for c,v in enumerate(row) if c < len(col_w)]
            lines.append("  │  ".join(padded))
            if i == 0:                       # separador tras encabezado
                lines.append("─" * min(sum(col_w) + 5*len(col_w), 90))
        return "\n".join(lines), None
    except Exception as e: return None, str(e)

# ──────────────────────────────────────────────
#  EXPORTAR
# ──────────────────────────────────────────────
def export_excel(data, path):
    if not XLSX_OK:
        messagebox.showerror("Error","pip install openpyxl"); return False
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Certificaciones"
    hdrs = ["ID","Nombre","Categoría","Fecha","Vencimiento","Tags","Descripción","Tipo","Ruta/URL"]
    from openpyxl.styles import Font, PatternFill, Alignment
    hdr_fill = PatternFill("solid", fgColor="4F46E5")
    hdr_font = Font(bold=True, color="FFFFFF")
    for col,h in enumerate(hdrs,1):
        cell = ws.cell(1,col,h)
        cell.font=hdr_font; cell.fill=hdr_fill
        cell.alignment=Alignment(horizontal="center")
    for r,c in enumerate(data,2):
        days  = days_to_expiry(c.get("fecha_venc",""))
        tags  = ", ".join(c.get("tags",[]))
        ruta  = c.get("url","") or c.get("ruta","")
        row   = [c.get("id",""),c.get("nombre",""),c.get("categoria",""),
                 c.get("fecha",""),c.get("fecha_venc",""),tags,
                 c.get("descripcion",""),c.get("tipo",""),ruta]
        for col,val in enumerate(row,1):
            ws.cell(r,col,val)
        # Color fila si vencido/próximo
        if days is not None:
            fill_color = "FEE2E2" if days<0 else ("FEF3C7" if days<=ALERT_DAYS else None)
            if fill_color:
                fill = PatternFill("solid",fgColor=fill_color)
                for col in range(1,len(hdrs)+1):
                    ws.cell(r,col).fill = fill
    for col in ws.columns:
        max_l = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_l+4,50)
    wb.save(path); return True

def export_pdf(data, path):
    if not RL_OK:
        messagebox.showerror("Error","pip install reportlab"); return False
    doc = SimpleDocTemplate(path, pagesize=A4,
                            leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title",parent=styles["Title"],
                                  fontSize=20,textColor=rl_colors.HexColor("#4f46e5"),
                                  spaceAfter=6)
    sub_style   = ParagraphStyle("sub",parent=styles["Normal"],
                                  fontSize=9,textColor=rl_colors.HexColor("#64748b"),
                                  spaceAfter=14)
    elements = [
        Paragraph("ORDERFOLDER", title_style),
        Paragraph(f"Reporte generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}  ·  Total: {len(data)} certificados", sub_style),
    ]
    hdrs = ["#","Nombre","Categoría","Fecha","Vencimiento","Tags"]
    tdata = [hdrs]
    for c in data:
        tags = ", ".join(c.get("tags",[])) or "—"
        tdata.append([str(c.get("id","")), c.get("nombre","")[:40],
                      c.get("categoria",""), c.get("fecha","") or "—",
                      c.get("fecha_venc","") or "—", tags])
    col_w = [1*cm,6.5*cm,3.5*cm,2.5*cm,2.5*cm,3*cm]
    t = Table(tdata, colWidths=col_w, repeatRows=1)
    ts = TableStyle([
        ("BACKGROUND",  (0,0),(-1,0),  rl_colors.HexColor("#4f46e5")),
        ("TEXTCOLOR",   (0,0),(-1,0),  rl_colors.white),
        ("FONTNAME",    (0,0),(-1,0),  "Helvetica-Bold"),
        ("FONTSIZE",    (0,0),(-1,0),  9),
        ("FONTSIZE",    (0,1),(-1,-1), 8),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[rl_colors.HexColor("#f8fafc"),
                                          rl_colors.white]),
        ("GRID",        (0,0),(-1,-1), 0.3, rl_colors.HexColor("#cbd5e1")),
        ("VALIGN",      (0,0),(-1,-1), "MIDDLE"),
        ("TOPPADDING",  (0,0),(-1,-1), 5),
        ("BOTTOMPADDING",(0,0),(-1,-1),5),
    ])
    # Filas vencidas/próximas en rojo/amarillo
    for i,c in enumerate(data,1):
        days = days_to_expiry(c.get("fecha_venc",""))
        if days is not None:
            if days < 0:
                ts.add("BACKGROUND",(0,i),(-1,i),rl_colors.HexColor("#fee2e2"))
            elif days <= ALERT_DAYS:
                ts.add("BACKGROUND",(0,i),(-1,i),rl_colors.HexColor("#fef3c7"))
    t.setStyle(ts)
    elements.append(t)
    doc.build(elements); return True

# ──────────────────────────────────────────────
#  UI HELPERS
# ──────────────────────────────────────────────
def apply_styles():
    s = ttk.Style(); s.theme_use("clam")
    s.configure("T.Treeview",background=C["card"],foreground=C["text2"],
                fieldbackground=C["card"],rowheight=32,font=fn(9),borderwidth=0)
    s.configure("T.Treeview.Heading",background=C["card2"],foreground=C["text"],
                font=fb(9),relief="flat",padding=(8,6))
    s.map("T.Treeview",background=[("selected",C["accent"])],
          foreground=[("selected","white")])
    s.configure("TScrollbar",background=C["card2"],troughcolor=C["card"],
                arrowcolor=C["sub"],borderwidth=0,relief="flat")
    s.configure("TCombobox",fieldbackground=C["card2"],background=C["card2"],
                foreground=C["text"],selectbackground=C["accent"],
                relief="flat",borderwidth=0)
    s.map("TCombobox",fieldbackground=[("readonly",C["card2"])],
          foreground=[("readonly",C["text"])])

def lh(h):
    r,g,b=int(h[1:3],16),int(h[3:5],16),int(h[5:7],16)
    return f"#{min(255,r+22):02x}{min(255,g+22):02x}{min(255,b+22):02x}"

def mk_btn(parent,text,bg,cmd,pad=(12,7),**kw):
    b=tk.Button(parent,text=text,font=fb(9),bg=bg,fg="white",
                activebackground=lh(bg),activeforeground="white",
                bd=0,padx=pad[0],pady=pad[1],cursor="hand2",
                relief="flat",command=cmd,**kw)
    b.bind("<Enter>",lambda e,_b=b,c=lh(bg):_b.config(bg=c))
    b.bind("<Leave>",lambda e,_b=b,c=bg:    _b.config(bg=c))
    return b

def mk_entry(parent,**kw):
    return tk.Entry(parent,bg=C["card2"],fg=C["text"],
                    insertbackground=C["cyan"],font=fn(10),bd=0,
                    highlightthickness=1,highlightbackground=C["border2"],
                    highlightcolor=C["accent"],relief="flat",**kw)

def mk_sep(parent):
    return tk.Frame(parent,bg=C["border"],height=1)

def mk_lbl(parent,text,size=10,bold=False,color=None,bg=None,**kw):
    return tk.Label(parent,text=text,font=fb(size) if bold else fn(size),
                    bg=bg or C["bg"],fg=color or C["text"],**kw)

def open_file(ruta):
    if not ruta:
        messagebox.showwarning("Sin archivo","Este registro no tiene archivo."); return
    if ruta.startswith("http"): webbrowser.open(ruta); return
    if not os.path.exists(ruta):
        messagebox.showerror("No encontrado",f"No encontrado:\n{ruta}"); return
    try:
        if sys.platform=="win32":    os.startfile(ruta)
        elif sys.platform=="darwin": subprocess.call(["open",ruta])
        else:                        subprocess.call(["xdg-open",ruta])
    except Exception as e: messagebox.showerror("Error",str(e))

def tipo_color(ext):
    return {".pdf":C["red"],".docx":C["accent"],".doc":C["accent"],
            ".xlsx":C["green"],".xls":C["green"],
            ".txt":C["orange"],".csv":C["yellow"],
            "url":C["cyan"]}.get(ext,C["purple"])

# ──────────────────────────────────────────────
#  PANEL PREVIEW
# ──────────────────────────────────────────────
class PreviewPanel(tk.Frame):
    def __init__(self,parent):
        super().__init__(parent,bg=C["card"],width=380)
        self.pack_propagate(False)
        self._img_ref=None; self._pdf_doc=None
        self._pdf_page=0; self._total_pg=1
        self._open_cmd=None; self._current=None
        self._build()

    def _build(self):
        hdr=tk.Frame(self,bg=C["card2"],height=44)
        hdr.pack(fill="x"); hdr.pack_propagate(False)
        self.lbl_tipo=tk.Label(hdr,text="PREVISUALIZACIÓN",
                                font=("Consolas",7,"bold"),bg=C["card2"],fg=C["sub"])
        self.lbl_tipo.pack(side="left",padx=14,pady=13)
        self.btn_open=mk_btn(hdr,"↗ Abrir",C["accent"],
                             lambda:self._open_cmd and self._open_cmd(),pad=(8,4))
        self.btn_open.pack(side="right",padx=10,pady=8)
        mk_sep(self).pack(fill="x")

        self.area=tk.Frame(self,bg=C["bg2"])
        self.area.pack(fill="both",expand=True)
        self.canvas=tk.Canvas(self.area,bg=C["bg2"],bd=0,highlightthickness=0)
        self.canvas.pack(fill="both",expand=True)
        self.canvas.bind("<Configure>",lambda e:self._redraw())

        txt_f=tk.Frame(self.area,bg=C["bg2"])
        self._txt_scroll=ttk.Scrollbar(txt_f)
        self.txt_box=tk.Text(txt_f,bg=C["card"],fg=C["text2"],
                              font=("Consolas",8),bd=0,wrap="none",
                              highlightthickness=0,relief="flat",
                              yscrollcommand=self._txt_scroll.set,state="disabled")
        self._txt_scroll.config(command=self.txt_box.yview)
        self._txt_scroll.pack(side="right",fill="y")
        self.txt_box.pack(fill="both",expand=True)
        self._txt_frame=txt_f

        self.pg_bar=tk.Frame(self,bg=C["card2"],height=36)
        self.pg_bar.pack(fill="x",side="bottom"); self.pg_bar.pack_propagate(False)
        mk_btn(self.pg_bar,"◀",C["card3"],self._prev,pad=(10,4)).pack(side="left",padx=6,pady=4)
        mk_btn(self.pg_bar,"▶",C["card3"],self._next,pad=(10,4)).pack(side="right",padx=6,pady=4)
        self.lbl_pg=tk.Label(self.pg_bar,text="",font=fn(9),bg=C["card2"],fg=C["sub"])
        self.lbl_pg.pack(expand=True)
        self.pg_bar.pack_forget()
        self._show_empty()

    def _show_canvas(self):
        self._txt_frame.pack_forget()
        self.canvas.pack(fill="both",expand=True)

    def _show_text(self):
        self.canvas.pack_forget()
        self._txt_frame.pack(fill="both",expand=True)

    def _show_empty(self):
        self._show_canvas(); self.canvas.delete("all")
        self.canvas.update_idletasks()
        w=self.canvas.winfo_width() or 340; h=self.canvas.winfo_height() or 420
        self.canvas.create_text(w//2,h//2,text="📄\n\nSelecciona un\ncertificado",
                                 fill=C["sub"],font=fn(11),justify="center")

    def load(self,cert,open_cmd):
        self._open_cmd=open_cmd; self._current=cert
        self.pg_bar.pack_forget()
        ruta=cert.get("ruta","") or ""; url=cert.get("url","") or ""
        tipo=cert.get("tipo","")
        ext=tipo or (os.path.splitext(ruta)[1].lower() if ruta else "url" if url else "")
        icon=TIPO_ICON.get(ext,"📄"); color=tipo_color(ext)
        self.lbl_tipo.config(text=f"{icon}  {ext.upper().lstrip('.') or 'URL'}",fg=color)
        if url and not ruta:  self._show_url(cert,url); return
        if not ruta or not os.path.exists(ruta): self._show_empty(); return
        ext2=os.path.splitext(ruta)[1].lower()
        if ext2==".pdf":               self._load_pdf(ruta)
        elif ext2 in EXTS_IMG:         self._load_img(ruta)
        elif ext2 in {".docx",".doc"}: self._load_docx(ruta)
        elif ext2 in {".xlsx",".xls"}: self._load_xlsx(ruta)
        elif ext2 == ".txt":           self._load_text_file(ruta, "txt")
        elif ext2 == ".csv":           self._load_text_file(ruta, "csv")
        else:                          self._show_empty()

    def _show_url(self,cert,url):
        self._show_text()
        self.txt_box.config(state="normal"); self.txt_box.delete("1.0","end")
        self.txt_box.insert("end",
            f"🔗  Google Sheets / URL\n\n"
            f"Nombre:    {cert.get('nombre','')}\n"
            f"Categoría: {cert.get('categoria','')}\n"
            f"Fecha:     {cert.get('fecha','')}\n\n"
            f"URL:\n{url}\n\n──────────────────\n"
            f"Clic en '↗ Abrir' para abrir en el navegador.")
        self.txt_box.config(state="disabled")

    def _load_img(self,ruta):
        self.pg_bar.pack_forget(); self._pdf_doc=None; self._show_canvas()
        self.canvas.update_idletasks()
        try:
            w=max(self.canvas.winfo_width(),340); h=max(self.canvas.winfo_height(),420)
            img=Image.open(ruta); img.thumbnail((w,h),Image.LANCZOS)
            self._img_ref=ImageTk.PhotoImage(img)
            self.canvas.delete("all")
            self.canvas.create_image(w//2,h//2,image=self._img_ref,anchor="center")
        except Exception as e: self._err(str(e))

    def _load_pdf(self,ruta):
        if not PDF_OK: self._msg("pip install pymupdf",C["yellow"]); return
        try:
            self._pdf_doc=fitz.open(ruta); self._total_pg=len(self._pdf_doc)
            self._pdf_page=0; self._render_pdf()
            if self._total_pg>1: self.pg_bar.pack(fill="x",side="bottom")
        except Exception as e: self._err(str(e))

    def _render_pdf(self):
        if not self._pdf_doc: return
        self._show_canvas(); self.canvas.update_idletasks()
        w=max(self.canvas.winfo_width(),340); h=max(self.canvas.winfo_height(),420)
        try:
            page=self._pdf_doc[self._pdf_page]; zoom=w/page.rect.width
            pix=page.get_pixmap(matrix=fitz.Matrix(zoom,zoom),alpha=False)
            img=Image.frombytes("RGB",[pix.width,pix.height],pix.samples)
            img.thumbnail((w,h),Image.LANCZOS)
            self._img_ref=ImageTk.PhotoImage(img)
            self.canvas.delete("all")
            self.canvas.create_image(w//2,h//2,image=self._img_ref,anchor="center")
            self.lbl_pg.config(text=f"Página {self._pdf_page+1} / {self._total_pg}")
        except Exception as e: self._err(str(e))

    def _load_docx(self,ruta):
        self.pg_bar.pack_forget(); self._pdf_doc=None
        content,err=read_docx(ruta)
        if err: self._msg(f"⚠ {err}",C["yellow"]); return
        self._show_text(); self.txt_box.config(state="normal")
        self.txt_box.delete("1.0","end"); self.txt_box.insert("end",content)
        self.txt_box.config(state="disabled")

    def _load_xlsx(self,ruta):
        self.pg_bar.pack_forget(); self._pdf_doc=None
        content,err=read_xlsx(ruta)
        if err: self._msg(f"⚠ {err}",C["yellow"]); return
        self._show_text(); self.txt_box.config(state="normal")
        self.txt_box.delete("1.0","end"); self.txt_box.insert("end",content)
        self.txt_box.config(state="disabled")

    def _load_text_file(self, ruta, kind):
        """Carga TXT (párrafos) o CSV (tabla alineada)."""
        self.pg_bar.pack_forget(); self._pdf_doc = None
        content, err = (read_txt(ruta) if kind=="txt" else read_csv(ruta))
        if err:
            self._msg(f"⚠ No se pudo leer:\n{err}", C["yellow"]); return
        self._show_text()
        # Color de encabezado distinto para CSV
        self.txt_box.config(state="normal",
                             font=("Consolas", 8) if kind=="csv" else ("Segoe UI", 9))
        self.txt_box.delete("1.0","end")
        if kind == "csv":
            lines = content.split("\n")
            # Primera línea = encabezado en color acento
            if lines:
                self.txt_box.insert("end", lines[0]+"\n", "header")
                self.txt_box.insert("end", "\n".join(lines[1:]))
            self.txt_box.tag_configure("header",
                foreground=C["cyan"], font=("Consolas",8,"bold"))
        else:
            self.txt_box.insert("end", content)
        self.txt_box.config(state="disabled")

    def _prev(self):
        if self._pdf_doc and self._pdf_page>0: self._pdf_page-=1; self._render_pdf()

    def _next(self):
        if self._pdf_doc and self._pdf_page<self._total_pg-1: self._pdf_page+=1; self._render_pdf()

    def _redraw(self):
        if not self._current: return
        ruta=self._current.get("ruta","")
        if not ruta: return
        ext=os.path.splitext(ruta)[1].lower()
        if ext==".pdf" and self._pdf_doc: self._render_pdf()
        elif ext in EXTS_IMG and os.path.exists(ruta): self._load_img(ruta)

    def _msg(self,text,color=None):
        self._show_canvas(); self.canvas.delete("all")
        w=self.canvas.winfo_width() or 340; h=self.canvas.winfo_height() or 420
        self.canvas.create_text(w//2,h//2,text=text,fill=color or C["sub"],
                                 font=fn(10),justify="center")

    def _err(self,msg): self._msg(f"⚠️ Error\n{msg[:80]}",C["red"])

    def clear(self):
        self._pdf_doc=None; self._current=None
        self.pg_bar.pack_forget(); self._show_empty()


# ──────────────────────────────────────────────
#  APP PRINCIPAL
# ──────────────────────────────────────────────
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("OrderFolder")
        self.geometry("1480x840"); self.minsize(1100,680)
        self.configure(bg=C["bg"])
        self.cfg  = load_cfg()
        # Aplicar tema guardado
        C.update(THEMES.get(self.cfg.get("theme","dark"), THEMES["dark"]))
        apply_styles()
        self.data = load_data()
        self._build()
        if self.cfg.get("carpeta_raiz") and os.path.isdir(self.cfg["carpeta_raiz"]):
            self._sync_silent()
        self._check_alerts()
        self._nav("lista", self.view_lista)

    # ── Layout ───────────────────────────────
    def _build(self):
        self.sb=tk.Frame(self,bg=C["sidebar"],width=240)
        self.sb.pack(side="left",fill="y"); self.sb.pack_propagate(False)
        self._build_sidebar()
        tk.Frame(self,bg=C["border"],width=1).pack(side="left",fill="y")
        self.content=tk.Frame(self,bg=C["bg"])
        self.content.pack(side="left",fill="both",expand=True)

    def _build_sidebar(self):
        logo=tk.Frame(self.sb,bg=C["sidebar"])
        logo.pack(fill="x",pady=(24,6))
        tk.Label(logo,text="ORDER",font=("Consolas",28,"bold"),
                 bg=C["sidebar"],fg=C["accent"]).pack()
        tk.Label(logo,text="F O L D E R",
                 font=("Consolas",7,"bold"),bg=C["sidebar"],fg=C["sub"]).pack()
        mk_sep(self.sb).pack(fill="x",padx=20,pady=12)

        # Carpeta activa
        cf=tk.Frame(self.sb,bg=C["card"],highlightthickness=1,
                    highlightbackground=C["border2"])
        cf.pack(fill="x",padx=14,pady=(0,8))
        tk.Label(cf,text="CARPETA ACTIVA",font=("Consolas",7,"bold"),
                 bg=C["card"],fg=C["sub"]).pack(anchor="w",padx=10,pady=(8,2))
        self.lbl_carpeta=tk.Label(cf,text=self._nom_carpeta(),font=fn(9),
                                   bg=C["card"],fg=C["cyan"],wraplength=190,justify="left")
        self.lbl_carpeta.pack(anchor="w",padx=10,pady=(0,6))
        mk_btn(cf,"⟳  Cambiar carpeta",C["accent"],
               self._cambiar_carpeta,pad=(10,5)).pack(fill="x",padx=10,pady=(0,10))

        mk_sep(self.sb).pack(fill="x",padx=20,pady=8)

        self.nav_btns={}
        for key,txt,cmd in [
            ("lista",  "  🗂   Lista",          self.view_lista),
            ("buscar", "  🔍   Buscar",          self.view_buscar),
            ("venc",   "  ⏰   Vencimientos",    self.view_venc),
            ("agreg",  "  ＋   Agregar",         self.view_agregar),
            ("export", "  ⬇   Exportar",         self.view_export),
            ("stats",  "  ◎   Estadísticas",    self.view_stats),
        ]:
            b=tk.Button(self.sb,text=txt,font=fn(10),bg=C["sidebar"],fg=C["text2"],
                        activebackground=C["nav_sel"],activeforeground=C["text"],
                        bd=0,padx=10,pady=12,anchor="w",cursor="hand2",relief="flat",
                        command=lambda c=cmd,k=key:self._nav(k,c))
            b.pack(fill="x",padx=8,pady=1); self.nav_btns[key]=b

        mk_sep(self.sb).pack(fill="x",padx=20,pady=8)

        # Leyenda tipos
        leg=tk.Frame(self.sb,bg=C["sidebar"])
        leg.pack(fill="x",padx=14)
        tk.Label(leg,text="TIPOS",font=("Consolas",7,"bold"),
                 bg=C["sidebar"],fg=C["sub"]).pack(anchor="w",pady=(0,4))
        for icon,label,color in [("📄","PDF",C["red"]),("🖼","Imagen",C["purple"]),
                                   ("📝","Word",C["accent"]),("📊","Excel",C["green"]),
                                   ("🔗","URL",C["cyan"])]:
            r=tk.Frame(leg,bg=C["sidebar"]); r.pack(fill="x",pady=1)
            tk.Label(r,text=icon,font=fn(9),bg=C["sidebar"],fg=color).pack(side="left")
            tk.Label(r,text=f" {label}",font=fn(8),bg=C["sidebar"],fg=C["sub"]).pack(side="left")

        # Toggle tema
        mk_sep(self.sb).pack(fill="x",padx=20,pady=8)
        theme_name = "dark" if C["bg"]=="#080b14" else "light"
        self.btn_theme=mk_btn(self.sb,
                               "☀  Modo claro" if theme_name=="dark" else "🌙  Modo oscuro",
                               C["card3"],self._toggle_theme,pad=(10,6))
        self.btn_theme.pack(fill="x",padx=14,pady=(0,8))

        mk_sep(self.sb).pack(fill="x",padx=20,side="bottom",pady=0)
        # Crédito sutil + contador
        credit_frame=tk.Frame(self.sb,bg=C["sidebar"])
        credit_frame.pack(side="bottom",pady=(4,6))
        tk.Label(credit_frame,text="Creado por Hack14",
                 font=("Consolas",7),bg=C["sidebar"],fg=C["border2"]).pack()
        self.lbl_tot=tk.Label(self.sb,text="",font=("Consolas",8),
                               bg=C["sidebar"],fg=C["sub"])
        self.lbl_tot.pack(side="bottom",pady=(6,2))
        self._upd_tot()

    def _nom_carpeta(self):
        r=self.cfg.get("carpeta_raiz","")
        return os.path.basename(r) if r else "No seleccionada"

    def _nav(self,key,cmd):
        for k,b in self.nav_btns.items():
            b.config(bg=C["nav_sel"] if k==key else C["sidebar"],
                     fg=C["text"] if k==key else C["text2"],
                     font=fb(10) if k==key else fn(10))
        cmd()

    def _clear(self):
        for w in self.content.winfo_children(): w.destroy()

    def _upd_tot(self):
        self.lbl_tot.config(text=f"▸  {len(self.data)} certificados")

    def _topbar(self,titulo,sub=""):
        bar=tk.Frame(self.content,bg=C["bg2"],height=56)
        bar.pack(fill="x"); bar.pack_propagate(False)
        inn=tk.Frame(bar,bg=C["bg2"])
        inn.pack(fill="both",expand=True,padx=28)
        tk.Label(inn,text=titulo,font=fb(13),bg=C["bg2"],fg=C["text"]).pack(side="left",pady=14)
        if sub: tk.Label(inn,text=sub,font=fn(9),bg=C["bg2"],fg=C["sub"]).pack(side="left",padx=10)
        mk_sep(self.content).pack(fill="x")
        return inn

    # ── Tema ─────────────────────────────────
    def _toggle_theme(self):
        current = "dark" if C["bg"]=="#080b14" else "light"
        new     = "light" if current=="dark" else "dark"
        C.update(THEMES[new])
        self.cfg["theme"]=new; save_cfg(self.cfg)
        apply_styles()
        # Reconstruir UI completa
        for w in self.winfo_children(): w.destroy()
        self.configure(bg=C["bg"])
        self._build()
        self._nav("lista",self.view_lista)

    # ── Sync ─────────────────────────────────
    def _cambiar_carpeta(self):
        nueva=filedialog.askdirectory(
            title="Selecciona tu carpeta de certificaciones",
            initialdir=self.cfg.get("carpeta_raiz") or os.path.expanduser("~"))
        if not nueva: return

        carpeta_anterior = self.cfg.get("carpeta_raiz","")
        es_carpeta_nueva = nueva != carpeta_anterior

        self.cfg["carpeta_raiz"]=nueva; save_cfg(self.cfg)
        self.lbl_carpeta.config(text=os.path.basename(nueva))

        # Si es una carpeta DISTINTA, preguntar si limpiar registros anteriores
        if es_carpeta_nueva and carpeta_anterior and self.data:
            resp = messagebox.askyesnocancel(
                "Nueva carpeta detectada",
                f"Estás cambiando de carpeta.\n\n"
                f"  Anterior: {os.path.basename(carpeta_anterior)}\n"
                f"  Nueva:    {os.path.basename(nueva)}\n\n"
                f"¿Deseas conservar los registros anteriores?\n\n"
                f"  ✅ Sí  → Mantiene todo y agrega los nuevos\n"
                f"  ❌ No  → Limpia la lista y carga solo la nueva carpeta\n"
                f"  Cancelar → Abortar cambio de carpeta")

            if resp is None:          # Cancelar → revertir
                self.cfg["carpeta_raiz"] = carpeta_anterior
                save_cfg(self.cfg)
                self.lbl_carpeta.config(text=os.path.basename(carpeta_anterior) if carpeta_anterior else "No seleccionada")
                return
            elif resp is False:       # No → limpiar registros anteriores
                self.data = []
                save_data(self.data)

        self._sync(True)

    def _sync_silent(self):
        try:
            sc=scan(self.cfg["carpeta_raiz"])
            self.data=sync(self.data,sc); save_data(self.data); self._upd_tot()
        except Exception as e: messagebox.showerror("Error",str(e))

    def _sync(self,show=False):
        try:
            raiz=self.cfg.get("carpeta_raiz",""); sc=scan(raiz)
            ant=len(self.data); self.data=sync(self.data,sc)
            save_data(self.data); self._upd_tot()
            if show:
                messagebox.showinfo("Sincronización",
                    f"📁 {os.path.basename(raiz)}\n"
                    f"📄 Leídos: {len(sc)}\n"
                    f"✅ Total:  {len(self.data)}\n"
                    f"➕ Nuevos: {max(0,len(self.data)-ant)}")
        except Exception as e: messagebox.showerror("Error",str(e))
        self._nav("lista",self.view_lista)

    # ── Alertas de vencimiento ────────────────
    def _check_alerts(self):
        alertas=[c for c in self.data
                 if days_to_expiry(c.get("fecha_venc","")) is not None
                 and (days_to_expiry(c.get("fecha_venc","")) or 999) <= ALERT_DAYS]
        if alertas:
            venc  =[c for c in alertas if (days_to_expiry(c.get("fecha_venc","")) or 0)<0]
            prox  =[c for c in alertas if (days_to_expiry(c.get("fecha_venc","")) or 0)>=0]
            msg   = f"⏰ Alertas de vencimiento\n\n"
            if venc:
                msg += f"🔴 Vencidos ({len(venc)}):\n"
                msg += "\n".join(f"  • {c['nombre']}" for c in venc[:5])
                msg += "\n\n"
            if prox:
                msg += f"🟡 Vencen en ≤{ALERT_DAYS} días ({len(prox)}):\n"
                msg += "\n".join(f"  • {c['nombre']}" for c in prox[:5])
            messagebox.showwarning("Certificados por vencer",msg)

    # ── Tabla compartida ─────────────────────
    def _mk_table(self,parent,cols,widths):
        wrap=tk.Frame(parent,bg=C["card"],highlightthickness=1,
                      highlightbackground=C["border2"])
        wrap.pack(fill="both",expand=True,padx=24,pady=(0,8))
        tree=ttk.Treeview(wrap,columns=cols,show="headings",
                           style="T.Treeview",selectmode="browse")
        for c,w in zip(cols,widths):
            tree.heading(c,text=c)
            tree.column(c,width=w,
                        anchor="center" if c in ("ID","Fecha","Tipo","✓","Estado") else "w",
                        stretch=(c=="Nombre"))
        vsb=ttk.Scrollbar(wrap,orient="vertical",command=tree.yview)
        hsb=ttk.Scrollbar(wrap,orient="horizontal",command=tree.xview)
        tree.configure(yscrollcommand=vsb.set,xscrollcommand=hsb.set)
        hsb.pack(side="bottom",fill="x"); vsb.pack(side="right",fill="y")
        tree.pack(fill="both",expand=True)
        tree.tag_configure("par",   background=C["card"])
        tree.tag_configure("impar", background=C["card2"])
        tree.tag_configure("miss",  background="#1a0d0d",foreground="#f87171")
        tree.tag_configure("expired",background="#2a0d0d",foreground="#fca5a5")
        tree.tag_configure("soon",  background="#2a1f00",foreground="#fcd34d")
        return tree

    def _fill(self,tree,lista):
        tree.delete(*tree.get_children())
        for i,c in enumerate(lista):
            ruta=c.get("ruta",""); url=c.get("url","")
            existe=os.path.exists(ruta) if ruta else bool(url)
            ext=c.get("tipo","") or (os.path.splitext(ruta)[1].lower() if ruta else "url" if url else "")
            icon=TIPO_ICON.get(ext,"📄")
            days=days_to_expiry(c.get("fecha_venc",""))
            status,_=expiry_status(days)
            tag=("expired" if status=="expired" else
                 "soon"    if status=="soon"    else
                 "miss"    if not existe        else
                 "par"     if i%2==0            else "impar")
            tree.insert("","end",iid=str(c["id"]),tags=(tag,),values=(
                c["id"],c.get("nombre",""),c.get("categoria",""),
                c.get("fecha",""),icon,
                ", ".join(c.get("tags",[])) or "—",
                "✓" if existe else "✗"))

    def _get_sel(self,tree):
        sel=tree.selection()
        if not sel: messagebox.showwarning("Sin selección","Selecciona un certificado."); return None
        try: return next((c for c in self.data if c["id"]==int(sel[0])),None)
        except: return None

    def _open_cert(self,cert):
        if not cert: return
        open_file(cert.get("url","") or cert.get("ruta",""))

    def _sel_preview(self,tree):
        cert=self._get_sel(tree)
        if cert and hasattr(self,"preview"):
            self.preview.load(cert,lambda c=cert:self._open_cert(c))

    # ── VISTA: Lista ─────────────────────────
    def view_lista(self):
        self._clear()
        bar=self._topbar("🗂  Lista completa",f"{len(self.data)} certificados")
        mk_btn(bar,"⟳  Re-escanear",C["cyan"],lambda:self._sync(True)).pack(side="right",pady=10)
        if not self.cfg.get("carpeta_raiz"): self._welcome(); return

        split=tk.Frame(self.content,bg=C["bg"])
        split.pack(fill="both",expand=True)
        left=tk.Frame(split,bg=C["bg"])
        left.pack(side="left",fill="both",expand=True)
        tk.Frame(split,bg=C["border"],width=1).pack(side="left",fill="y")
        self.preview=PreviewPanel(split); self.preview.pack(side="right",fill="y")

        cols=("ID","Nombre","Categoría","Fecha","Tipo","Tags","✓")
        widths=[40,220,120,90,44,160,36]
        self.tree_lista=self._mk_table(left,cols,widths)
        self._fill(self.tree_lista,self.data)
        self.tree_lista.bind("<<TreeviewSelect>>",lambda e:self._sel_preview(self.tree_lista))

        f=tk.Frame(left,bg=C["bg"]); f.pack(fill="x",padx=24,pady=(0,14))
        mk_btn(f,"↗  Abrir",     C["cyan"],   lambda:self._open_cert(self._get_sel(self.tree_lista))).pack(side="left",padx=3)
        mk_btn(f,"✏️  Editar",    C["accent"], lambda:self._edit(self.tree_lista)).pack(side="left",padx=3)
        mk_btn(f,"🗑  Eliminar",  C["red"],    lambda:self._delete(self.tree_lista)).pack(side="left",padx=3)

    def _welcome(self):
        f=tk.Frame(self.content,bg=C["bg"]); f.pack(expand=True)
        tk.Label(f,text="OF",font=("Consolas",52,"bold"),bg=C["bg"],fg=C["accent"]).pack(pady=(50,4))
        mk_lbl(f,"O R D E R F O L D E R",11,bold=True,color=C["sub"]).pack()
        mk_sep(f).pack(fill="x",pady=20,padx=60)
        mk_lbl(f,"Selecciona tu carpeta de certificaciones para comenzar.",10,color=C["text2"]).pack()
        mk_lbl(f,"Detecta PDF, Imágenes, Word y Excel automáticamente.",9,color=C["sub"]).pack(pady=4)
        mk_btn(f,"📂  Seleccionar carpeta",C["accent"],self._cambiar_carpeta,pad=(18,10)).pack(pady=20)

    # ── VISTA: Buscar ────────────────────────
    def view_buscar(self):
        self._clear(); self._topbar("🔍  Buscar / Filtrar")
        cats=sorted({c.get("categoria","") for c in self.data if c.get("categoria","").strip()})
        all_tags=sorted({t for c in self.data for t in c.get("tags",[])})
        tipos=["Todos","PDF","Imagen","Word","Excel","URL"]

        pf=tk.Frame(self.content,bg=C["card2"],highlightthickness=1,highlightbackground=C["border2"])
        pf.pack(fill="x",padx=24,pady=(10,8))
        inn=tk.Frame(pf,bg=C["card2"]); inn.pack(fill="x",padx=14,pady=12)

        for col,txt in enumerate(["TEXTO","CATEGORÍA","TIPO","TAG"]):
            tk.Label(inn,text=txt,font=("Consolas",7,"bold"),
                     bg=C["card2"],fg=C["sub"]).grid(row=0,column=col,sticky="w",padx=(0,4))

        self.v_txt=tk.StringVar()
        e=mk_entry(inn,textvariable=self.v_txt,width=24)
        e.grid(row=1,column=0,padx=(0,12),sticky="ew",ipady=4)

        self.v_cat=tk.StringVar(value="Todas")
        ttk.Combobox(inn,textvariable=self.v_cat,values=["Todas"]+cats,
                     state="readonly",width=18,font=fn(10)).grid(row=1,column=1,padx=(0,12),sticky="ew",ipady=4)

        self.v_tipo=tk.StringVar(value="Todos")
        ttk.Combobox(inn,textvariable=self.v_tipo,values=tipos,
                     state="readonly",width=10,font=fn(10)).grid(row=1,column=2,padx=(0,12),sticky="ew",ipady=4)

        self.v_tag=tk.StringVar(value="Todos")
        ttk.Combobox(inn,textvariable=self.v_tag,values=["Todos"]+all_tags,
                     state="readonly",width=14,font=fn(10)).grid(row=1,column=3,padx=(0,12),sticky="ew",ipady=4)

        bf=tk.Frame(inn,bg=C["card2"]); bf.grid(row=1,column=4)
        mk_btn(bf,"Buscar", C["accent"],self._do_search, pad=(14,6)).pack(side="left",padx=3)
        mk_btn(bf,"Limpiar",C["card3"], self._clear_srch,pad=(10,6)).pack(side="left",padx=3)
        inn.columnconfigure(0,weight=1)

        split=tk.Frame(self.content,bg=C["bg"]); split.pack(fill="both",expand=True)
        left=tk.Frame(split,bg=C["bg"]); left.pack(side="left",fill="both",expand=True)
        tk.Frame(split,bg=C["border"],width=1).pack(side="left",fill="y")
        self.preview=PreviewPanel(split); self.preview.pack(side="right",fill="y")

        cols=("ID","Nombre","Categoría","Fecha","Tipo","Tags","✓")
        widths=[40,220,120,90,44,160,36]
        self.tree_buscar=self._mk_table(left,cols,widths)
        self._fill(self.tree_buscar,self.data)
        self.tree_buscar.bind("<<TreeviewSelect>>",lambda e:self._sel_preview(self.tree_buscar))

        self.lbl_res=tk.Label(left,text=f"{len(self.data)} resultados",
                               font=("Consolas",8),bg=C["bg"],fg=C["sub"])
        self.lbl_res.pack(anchor="e",padx=26,pady=(0,4))
        f=tk.Frame(left,bg=C["bg"]); f.pack(fill="x",padx=24,pady=(0,14))
        mk_btn(f,"↗  Abrir",C["cyan"],lambda:self._open_cert(self._get_sel(self.tree_buscar))).pack(side="left",padx=3)
        e.bind("<Return>",lambda _:self._do_search()); e.focus_set()

    def _do_search(self):
        t=self.v_txt.get().lower().strip(); cat=self.v_cat.get()
        tipo=self.v_tipo.get(); tag=self.v_tag.get()
        tipo_map={"PDF":[".pdf"],"Imagen":list(EXTS_IMG),
                  "Word":[".docx",".doc"],"Excel":[".xlsx",".xls"],"URL":["url"]}
        exts=tipo_map.get(tipo,[])
        def match_tipo(c):
            if tipo=="Todos": return True
            ruta=c.get("ruta",""); url=c.get("url","")
            ext=c.get("tipo","") or (os.path.splitext(ruta)[1].lower() if ruta else "url" if url else "")
            if tipo=="URL": return bool(url and not ruta)
            return ext in exts
        res=[c for c in self.data
             if (not t or t in c.get("nombre","").lower()
                       or t in c.get("descripcion","").lower()
                       or t in c.get("categoria","").lower())
             and (cat=="Todas" or c.get("categoria","")==cat)
             and match_tipo(c)
             and (tag=="Todos" or tag in c.get("tags",[]))]
        self._fill(self.tree_buscar,res)
        n=len(res); self.lbl_res.config(text=f"{n} resultado{'s' if n!=1 else ''}")

    def _clear_srch(self):
        self.v_txt.set(""); self.v_cat.set("Todas")
        self.v_tipo.set("Todos"); self.v_tag.set("Todos")
        self._fill(self.tree_buscar,self.data)
        self.lbl_res.config(text=f"{len(self.data)} resultados")
        if hasattr(self,"preview"): self.preview.clear()

    # ── VISTA: Vencimientos ──────────────────
    def view_venc(self):
        self._clear()
        self._topbar("⏰  Vencimientos","Certificados con fecha de vencimiento")

        con_venc=[c for c in self.data if c.get("fecha_venc","").strip()]
        vencidos =[c for c in con_venc if (days_to_expiry(c.get("fecha_venc","")) or 1)<0]
        proximos =[c for c in con_venc if 0<=(days_to_expiry(c.get("fecha_venc","")) or 999)<=ALERT_DAYS]
        vigentes =[c for c in con_venc if (days_to_expiry(c.get("fecha_venc","")) or 999)>ALERT_DAYS]

        # Resumen
        row=tk.Frame(self.content,bg=C["bg"]); row.pack(fill="x",padx=24,pady=(14,16))
        for ico,txt,val,col in [("🔴","Vencidos",str(len(vencidos)),C["red"]),
                                  ("🟡",f"En ≤{ALERT_DAYS}d",str(len(proximos)),C["yellow"]),
                                  ("🟢","Vigentes",str(len(vigentes)),C["green"]),
                                  ("📅","Total con venc.",str(len(con_venc)),C["cyan"])]:
            f=tk.Frame(row,bg=C["card"],highlightthickness=1,highlightbackground=col)
            f.pack(side="left",expand=True,fill="x",padx=5)
            tk.Label(f,text=ico,font=fn(18),bg=C["card"],fg=col).pack(pady=(12,2))
            tk.Label(f,text=txt,font=fn(8),bg=C["card"],fg=C["sub"]).pack()
            tk.Label(f,text=val,font=("Consolas",22,"bold"),bg=C["card"],fg=col).pack(pady=(2,12))

        mk_sep(self.content).pack(fill="x",padx=24,pady=(0,10))

        cols=("ID","Nombre","Categoría","Fecha obtención","Vence","Días restantes","✓")
        widths=[40,220,120,110,110,110,36]
        wrap=tk.Frame(self.content,bg=C["card"],highlightthickness=1,
                      highlightbackground=C["border2"])
        wrap.pack(fill="both",expand=True,padx=24,pady=(0,14))
        tree=ttk.Treeview(wrap,columns=cols,show="headings",
                           style="T.Treeview",selectmode="browse")
        for c,w in zip(cols,widths):
            tree.heading(c,text=c); tree.column(c,width=w,
                anchor="center" if c in ("ID","Días restantes","✓") else "w")
        vsb=ttk.Scrollbar(wrap,orient="vertical",command=tree.yview)
        vsb.pack(side="right",fill="y"); tree.pack(fill="both",expand=True)
        tree.configure(yscrollcommand=vsb.set)
        tree.tag_configure("expired",background="#2a0d0d",foreground="#fca5a5")
        tree.tag_configure("soon",   background="#2a1f00",foreground="#fcd34d")
        tree.tag_configure("ok",     background=C["card"])

        ordered = sorted(con_venc,
                         key=lambda c: days_to_expiry(c.get("fecha_venc","")) or 9999)
        for c in ordered:
            days=days_to_expiry(c.get("fecha_venc",""))
            status,label=expiry_status(days)
            ruta=c.get("ruta",""); url=c.get("url","")
            existe=os.path.exists(ruta) if ruta else bool(url)
            tree.insert("","end",iid=str(c["id"]),tags=(status,),values=(
                c["id"],c.get("nombre",""),c.get("categoria",""),
                c.get("fecha",""),c.get("fecha_venc",""),label,"✓" if existe else "✗"))

    # ── VISTA: Exportar ──────────────────────
    def view_export(self):
        self._clear()
        bar=self._topbar("⬇  Exportar","Genera reportes de tu colección")

        f=tk.Frame(self.content,bg=C["bg"]); f.pack(expand=True)
        tk.Label(f,text="⬇",font=fn(36),bg=C["bg"],fg=C["accent"]).pack(pady=(40,8))
        mk_lbl(f,"Exportar reporte",14,bold=True).pack()
        mk_lbl(f,f"Se exportarán {len(self.data)} certificados",9,color=C["sub"]).pack(pady=4)
        mk_sep(f).pack(fill="x",pady=20,padx=60)

        cards=tk.Frame(f,bg=C["bg"]); cards.pack()

        # Excel
        xf=tk.Frame(cards,bg=C["card"],highlightthickness=1,highlightbackground=C["green"])
        xf.pack(side="left",padx=16,pady=8,ipadx=20,ipady=16)
        tk.Label(xf,text="📊",font=fn(28),bg=C["card"],fg=C["green"]).pack(pady=(12,4))
        mk_lbl(xf,"Exportar a Excel",10,bold=True,bg=C["card"]).pack()
        mk_lbl(xf,"Tabla completa con colores\npor estado de vencimiento",
               8,color=C["sub"],bg=C["card"]).pack(pady=4)
        mk_btn(xf,"📊  Guardar .xlsx",C["green"],self._export_xlsx,pad=(16,8)).pack(pady=(8,12))

        # PDF
        pf=tk.Frame(cards,bg=C["card"],highlightthickness=1,highlightbackground=C["red"])
        pf.pack(side="left",padx=16,pady=8,ipadx=20,ipady=16)
        tk.Label(pf,text="📄",font=fn(28),bg=C["card"],fg=C["red"]).pack(pady=(12,4))
        mk_lbl(pf,"Exportar a PDF",10,bold=True,bg=C["card"]).pack()
        mk_lbl(pf,"Reporte visual con tabla\ny resumen de vencimientos",
               8,color=C["sub"],bg=C["card"]).pack(pady=4)
        mk_btn(pf,"📄  Guardar .pdf",C["red"],self._export_pdf_rep,pad=(16,8)).pack(pady=(8,12))

    def _export_xlsx(self):
        path=filedialog.asksaveasfilename(
            title="Guardar Excel",defaultextension=".xlsx",
            initialfile=f"orderfolder_{datetime.now().strftime('%Y%m%d')}.xlsx",
            filetypes=[("Excel","*.xlsx")])
        if path and export_excel(self.data,path):
            messagebox.showinfo("✅ Exportado",f"Excel guardado en:\n{path}")
            open_file(path)

    def _export_pdf_rep(self):
        path=filedialog.asksaveasfilename(
            title="Guardar PDF",defaultextension=".pdf",
            initialfile=f"orderfolder_{datetime.now().strftime('%Y%m%d')}.pdf",
            filetypes=[("PDF","*.pdf")])
        if path and export_pdf(self.data,path):
            messagebox.showinfo("✅ Exportado",f"PDF guardado en:\n{path}")
            open_file(path)

    # ── VISTA: Estadísticas ──────────────────
    def view_stats(self):
        self._clear()
        self._topbar("◎  Estadísticas","Resumen de tu colección")
        if not self.data:
            mk_lbl(self.content,"Sin datos aún.",10,color=C["sub"]).pack(pady=60); return

        cats=Counter(c.get("categoria","?") for c in self.data)
        total=len(self.data); top=cats.most_common(1)[0]
        con_f=sum(1 for c in self.data if c.get("fecha","").strip())
        venc_cnt=sum(1 for c in self.data if c.get("fecha_venc","").strip())

        row=tk.Frame(self.content,bg=C["bg"]); row.pack(fill="x",padx=24,pady=(14,16))
        for ico,tag,sub,val,col in [
            ("📄","TOTAL",   "Certificados",str(total),   C["accent"]),
            ("🗂","CATS",    "Categorías",  str(len(cats)),C["cyan"]),
            ("🏆","TOP",     top[0],        str(top[1]),   C["green"]),
            ("⏰","VENC",    "Con vencim.", str(venc_cnt), C["yellow"]),
        ]:
            f=tk.Frame(row,bg=C["card"],highlightthickness=1,highlightbackground=col)
            f.pack(side="left",expand=True,fill="x",padx=5)
            tk.Label(f,text=ico,font=fn(20),bg=C["card"],fg=col).pack(pady=(14,2))
            tk.Label(f,text=tag,font=("Consolas",7,"bold"),bg=C["card"],fg=C["sub"]).pack()
            tk.Label(f,text=sub,font=fn(8),bg=C["card"],fg=C["sub"]).pack()
            tk.Label(f,text=val,font=("Consolas",22,"bold"),bg=C["card"],fg=col).pack(pady=(4,14))

        plt.rcParams.update({
            "figure.facecolor":C["bg2"],"axes.facecolor":C["card"],
            "axes.edgecolor":C["border2"],"text.color":C["text"],
            "xtick.color":C["sub"],"ytick.color":C["text2"],
            "grid.color":C["border"],"grid.linestyle":":","grid.alpha":0.5,
        })
        fig=plt.figure(figsize=(14,4.8),facecolor=C["bg2"])
        fig.subplots_adjust(left=0.03,right=0.97,top=0.88,bottom=0.24,wspace=0.4)
        ax1=fig.add_subplot(1,3,(1,2)); ax2=fig.add_subplot(1,3,3)

        etiq=list(cats.keys()); vals=list(cats.values())
        idx=sorted(range(len(vals)),key=lambda i:vals[i])
        etiq=[etiq[i] for i in idx]; vals=[vals[i] for i in idx]
        norm=plt.Normalize(min(vals),max(vals) if max(vals)>0 else 1)
        colors=plt.cm.cool(norm(vals))

        bars=ax1.barh(etiq,vals,color=colors,height=0.62,zorder=3)
        ax1.xaxis.grid(True,zorder=0); ax1.set_axisbelow(True)
        ax1.set_xlabel("Certificados",fontsize=8,labelpad=8)
        ax1.set_title("Por categoría",fontsize=10,color=C["text"],pad=10,fontweight="bold")
        ax1.tick_params(labelsize=8)
        for bar,v in zip(bars,vals):
            ax1.text(bar.get_width()+0.12,bar.get_y()+bar.get_height()/2,
                     str(v),va="center",ha="left",fontsize=8,color=C["text"],fontweight="bold")
        ax1.spines["top"].set_visible(False); ax1.spines["right"].set_visible(False)

        def get_tipo(c):
            ruta=c.get("ruta",""); url=c.get("url","")
            ext=c.get("tipo","") or (os.path.splitext(ruta)[1].lower() if ruta else "url" if url else "?")
            if ext in EXTS_IMG: return "Imagen"
            return {".pdf":"PDF",".docx":"Word",".doc":"Word",
                    ".xlsx":"Excel",".xls":"Excel","url":"URL"}.get(ext,ext.upper())
        tipos_cnt=Counter(get_tipo(c) for c in self.data)
        t_labels=list(tipos_cnt.keys()); t_vals=list(tipos_cnt.values())
        t_colors=[tipo_color({
            "PDF":".pdf","Word":".docx","Excel":".xlsx",
            "Imagen":".png","URL":"url"}.get(l,".pdf")) for l in t_labels]
        _,_,autotexts=ax2.pie(t_vals,labels=t_labels,colors=t_colors,
            autopct=lambda p:f"{p:.1f}%" if p>5 else "",startangle=90,
            pctdistance=0.75,
            wedgeprops={"width":0.55,"edgecolor":C["bg2"],"linewidth":2},
            textprops={"fontsize":8,"color":C["text2"]})
        for at in autotexts:
            at.set_color("white"); at.set_fontsize(8); at.set_fontweight("bold")
        ax2.text(0,0,f"{total}\ntotal",ha="center",va="center",
                 fontsize=10,color=C["text"],fontweight="bold")
        ax2.set_title("Por tipo de archivo",fontsize=10,color=C["text"],pad=10,fontweight="bold")

        canvas=FigureCanvasTkAgg(fig,master=self.content)
        canvas.draw()
        canvas.get_tk_widget().configure(bg=C["bg"],highlightthickness=0)
        canvas.get_tk_widget().pack(fill="both",expand=True,padx=24,pady=(0,16))

    # ── Acciones ─────────────────────────────
    def _edit(self,tree):
        c=self._get_sel(tree)
        if c: FormCert(self,cert=c,on_save=self._saved,
                       carpeta_raiz=self.cfg.get("carpeta_raiz",""))

    def _delete(self,tree):
        c=self._get_sel(tree)
        if not c: return
        if not messagebox.askyesno("Confirmar",
                f"¿Eliminar «{c['nombre']}»?\n⚠ El archivo original NO se borrará."): return
        self.data=[x for x in self.data if x["id"]!=c["id"]]
        save_data(self.data); self._upd_tot()
        self._fill(self.tree_lista,self.data)
        if hasattr(self,"preview"): self.preview.clear()

    def _saved(self,cert):
        ids=[c["id"] for c in self.data]
        if cert["id"] in ids:
            self.data=[cert if c["id"]==cert["id"] else c for c in self.data]
        else:
            cert["id"]=(max(ids)+1) if ids else 1; self.data.append(cert)
        save_data(self.data); self._upd_tot()
        self._nav("lista",self.view_lista)

    def view_agregar(self):
        FormCert(self,cert=None,on_save=self._saved,
                 carpeta_raiz=self.cfg.get("carpeta_raiz",""))


# ──────────────────────────────────────────────
#  FORMULARIO
# ──────────────────────────────────────────────
class FormCert(tk.Toplevel):
    def __init__(self,master,cert,on_save,carpeta_raiz=""):
        super().__init__(master)
        self.on_save=on_save; self.cert=cert or {}
        self.raiz=carpeta_raiz; self.ruta_n=None
        self.title("Editar" if cert else "Nuevo certificado")
        self.geometry("600x680"); self.configure(bg=C["bg"])
        self.resizable(False,False); self.grab_set(); self.focus_set()
        self._build()

    def _build(self):
        hdr=tk.Frame(self,bg=C["card2"],height=50)
        hdr.pack(fill="x"); hdr.pack_propagate(False)
        tk.Label(hdr,text="✏️  EDITAR" if self.cert else "＋  NUEVO CERTIFICADO",
                 font=("Consolas",10,"bold"),bg=C["card2"],fg=C["text"]).pack(side="left",padx=20,pady=14)
        mk_sep(self).pack(fill="x")

        # Scrollable body
        outer=tk.Frame(self,bg=C["bg"]); outer.pack(fill="both",expand=True)
        canvas=tk.Canvas(outer,bg=C["bg"],bd=0,highlightthickness=0)
        scroll=ttk.Scrollbar(outer,orient="vertical",command=canvas.yview)
        body=tk.Frame(canvas,bg=C["bg"])
        body.bind("<Configure>",lambda e:canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0,0),window=body,anchor="nw")
        canvas.configure(yscrollcommand=scroll.set)
        scroll.pack(side="right",fill="y"); canvas.pack(side="left",fill="both",expand=True)

        def lf(txt,row):
            tk.Label(body,text=txt,font=("Consolas",7,"bold"),
                     bg=C["bg"],fg=C["sub"]).grid(row=row,column=0,sticky="w",
                     padx=26,pady=(10,2))

        def ef(row,val="",col_span=2):
            e=mk_entry(body); e.insert(0,val)
            e.grid(row=row,column=0,columnspan=col_span,sticky="ew",padx=26,ipady=4)
            return e

        lf("NOMBRE *",0); self.e_nom=ef(1,self.cert.get("nombre",""))

        lf("CATEGORÍA *",2)
        existing=sorted({c.get("categoria","") for c in load_data() if c.get("categoria","").strip()})
        all_cats=sorted(set(existing)|{"Alcaldia Bogota","Campus","Colegiatura","Colsubsidio",
                        "Coursera","Fundacion Carlos Slim","Google","Linkedin","Platzi",
                        "Reconocimientos","Sena","Servinformacion","Solo Learnerd",
                        "Todo Code","Udemy","Universidad","Otra"})
        self.v_cat=tk.StringVar(value=self.cert.get("categoria",all_cats[0] if all_cats else "Otra"))
        ttk.Combobox(body,textvariable=self.v_cat,values=all_cats,
                     font=fn(10)).grid(row=3,column=0,columnspan=2,sticky="ew",padx=26,ipady=4)

        lf("FECHA OBTENCIÓN (YYYY-MM-DD)",4)
        self.e_fecha=ef(5,self.cert.get("fecha",datetime.today().strftime("%Y-%m-%d")))

        lf("FECHA VENCIMIENTO (YYYY-MM-DD)  — opcional",6)
        self.e_venc=ef(7,self.cert.get("fecha_venc",""))
        tk.Label(body,text="⏰ Si la completas recibirás alertas 30 días antes",
                 font=fn(7),bg=C["bg"],fg=C["yellow"]).grid(row=8,column=0,columnspan=2,
                 sticky="w",padx=26)

        lf("TAGS  (separados por coma)",9)
        self.e_tags=ef(10,", ".join(self.cert.get("tags",[])))
        tk.Label(body,text="Ej: python, backend, cloud",
                 font=fn(7),bg=C["bg"],fg=C["sub"]).grid(row=11,column=0,columnspan=2,
                 sticky="w",padx=26)

        lf("DESCRIPCIÓN",12); self.e_desc=ef(13,self.cert.get("descripcion",""))

        lf("ARCHIVO LOCAL  (PDF · Imagen · Word · Excel)",14)
        fa=tk.Frame(body,bg=C["bg"])
        fa.grid(row=15,column=0,columnspan=2,sticky="ew",padx=26,pady=(2,0))
        ruta_act=self.cert.get("ruta","") or self.cert.get("archivo","")
        self.lbl_arch=tk.Label(fa,text=os.path.basename(ruta_act) if ruta_act else "Sin archivo",
                                font=fn(8),bg=C["bg"],fg=C["sub"])
        self.lbl_arch.pack(side="left",expand=True,anchor="w")
        bf=tk.Frame(fa,bg=C["bg"]); bf.pack(side="right")
        mk_btn(bf,"📂 Apuntar",     C["cyan"],   self._pick,      pad=(8,4)).pack(side="left",padx=2)
        mk_btn(bf,"💾 Guardar copia",C["accent"], self._copy_save, pad=(8,4)).pack(side="left",padx=2)

        lf("URL  (Google Sheets u otro enlace)",16)
        self.e_url=ef(17,self.cert.get("url",""))
        tk.Label(body,text="Ej: https://docs.google.com/spreadsheets/...",
                 font=fn(7),bg=C["bg"],fg=C["sub"]).grid(row=18,column=0,columnspan=2,
                 sticky="w",padx=26,pady=(0,14))

        body.columnconfigure(0,weight=1)
        mk_sep(self).pack(fill="x")

        foot=tk.Frame(self,bg=C["card2"],height=52)
        foot.pack(fill="x"); foot.pack_propagate(False)
        mk_btn(foot,"💾  Guardar",C["green"], self._save,   pad=(16,8)).pack(side="right",padx=16,pady=10)
        mk_btn(foot,"Cancelar",   C["card3"], self.destroy, pad=(12,8)).pack(side="right",padx=4, pady=10)

    def _pick(self):
        p=filedialog.askopenfilename(
            title="Selecciona el archivo",
            initialdir=self.raiz or os.path.expanduser("~"),
            filetypes=[
                ("Todos los soportados",
                 "*.pdf *.png *.jpg *.jpeg *.webp *.docx *.xlsx *.xls *.txt *.csv"),
                ("PDF","*.pdf"),
                ("Imagen","*.png *.jpg *.jpeg *.webp"),
                ("Word","*.docx"),
                ("Excel","*.xlsx *.xls"),
                ("Texto","*.txt"),
                ("CSV","*.csv"),
                ("Todos","*.*"),
            ])
        if p:
            self.ruta_n=p
            self.lbl_arch.config(text=os.path.basename(p),
                                  fg=tipo_color(os.path.splitext(p)[1].lower()))

    def _copy_save(self):
        p=filedialog.askopenfilename(
            title="Selecciona el archivo a copiar",
            initialdir=self.raiz or os.path.expanduser("~"),
            filetypes=[
                ("Todos los soportados",
                 "*.pdf *.png *.jpg *.jpeg *.webp *.docx *.xlsx *.xls *.txt *.csv"),
                ("Todos","*.*"),
            ])
        if not p: return
        ext=os.path.splitext(p)[1]; nom=self.e_nom.get().strip() or "certificado"
        safe="".join(ch for ch in nom if ch.isalnum() or ch in " _-")
        dest=filedialog.asksaveasfilename(
            title="Guardar copia como...",
            initialdir=self.raiz or os.path.expanduser("~"),
            initialfile=f"{safe}{ext}",defaultextension=ext,
            filetypes=[("Mismo formato","*"+ext),("Todos","*.*")])
        if dest:
            shutil.copy2(p,dest); self.ruta_n=dest
            self.lbl_arch.config(text=os.path.basename(dest),fg=C["green"])
            messagebox.showinfo("Copia guardada",f"Guardado en:\n{dest}")

    def _save(self):
        nom=self.e_nom.get().strip()
        if not nom: messagebox.showwarning("Requerido","El nombre es obligatorio."); return
        tags=[t.strip() for t in self.e_tags.get().split(",") if t.strip()]
        ruta=self.ruta_n or self.cert.get("ruta","")
        ext=os.path.splitext(ruta)[1].lower() if ruta else ""
        self.on_save({
            "id":         self.cert.get("id",0),
            "nombre":     nom,
            "categoria":  self.v_cat.get(),
            "fecha":      self.e_fecha.get().strip(),
            "fecha_venc": self.e_venc.get().strip(),
            "descripcion":self.e_desc.get().strip(),
            "tags":       tags,
            "ruta":       ruta,
            "archivo":    self.cert.get("archivo",""),
            "tipo":       ext,
            "url":        self.e_url.get().strip(),
        })
        self.destroy()


if __name__ == "__main__":
    app = App()
    app.mainloop()